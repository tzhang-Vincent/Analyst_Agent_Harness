from agent.schema import Fact, Provenance
from agent.table import (
    build_rows,
    collect_periods,
    flatten,
    fmt_compact,
    period_labels_map,
    to_csv,
    to_excel,
)


def _fact(ticker, metric, end, label, value, accn):
    return Fact(
        metric=metric,
        period=end,
        period_label=label,
        provenance=Provenance(
            ticker=ticker,
            cik="0001045810",
            entity_name="NVIDIA CORP",
            xbrl_tag="Revenues",
            xbrl_label="Revenues",
            unit="USD",
            value=value,
            start_date="2025-04-28",
            end_date=end,
            fiscal_year=2026,
            fiscal_period="Q2",
            form="10-Q",
            accession_number=accn,
            filed_date="2025-08-26",
            source_url=f"https://www.sec.gov/Archives/edgar/data/1045810/x/{accn}-index.htm",
        ),
    )


def _facts():
    return {
        "NVDA": {
            "revenue": [
                _fact("NVDA", "revenue", "2025-07-27", "Q2 FY2026", 46743000000, "a1"),
                _fact("NVDA", "revenue", "2025-10-26", "Q3 FY2026", 57006000000, "a2"),
            ],
            "net_income": [
                _fact("NVDA", "net_income", "2025-07-27", "Q2 FY2026", 26422000000, "a3"),
            ],
        },
    }


def test_collect_periods_sorted():
    assert collect_periods(_facts()) == ["2025-07-27", "2025-10-26"]


def test_period_labels_map():
    labels = period_labels_map(_facts(), ["2025-07-27", "2025-10-26"])
    assert labels == {"2025-07-27": "Q2 FY2026", "2025-10-26": "Q3 FY2026"}


def test_build_rows():
    rows = build_rows(_facts(), ["2025-07-27", "2025-10-26"])
    assert len(rows) == 2  # NVDA/revenue and NVDA/net_income
    rev = next(r for r in rows if r["metric"] == "revenue")
    assert rev["cells"]["2025-10-26"].provenance.value == 57006000000
    assert "2025-10-26" not in rows[1]["cells"]  # net_income has no Q3


def test_fmt_compact():
    assert fmt_compact(46743000000) == "$46.74B"
    assert fmt_compact(2_500_000) == "$2.50M"
    assert fmt_compact(3_000_000_000_000) == "$3.00T"


def test_to_excel_two_sheets_and_hyperlink(tmp_path):
    path = tmp_path / "out.xlsx"
    to_excel(_facts(), str(path))
    from openpyxl import load_workbook

    wb = load_workbook(path)
    assert wb.sheetnames == ["Comparison", "Provenance"]
    # Comparison sheet has a hyperlinked cell (revenue Q3).
    ws = wb["Comparison"]
    assert ws.cell(row=1, column=1).value == "Company"
    # Provenance sheet has one row per fact (3 facts + header).
    pv = wb["Provenance"]
    assert pv.max_row == 4
    assert pv.cell(row=2, column=1).value == "NVDA"
    assert pv.cell(row=2, column=8).value == "Revenues"


def test_to_csv(tmp_path):
    path = tmp_path / "out.csv"
    facts = flatten(_facts())
    to_csv(facts, str(path))
    text = path.read_text()
    assert "ticker,entity,metric" in text
    assert "NVDA" in text
    assert "a2" in text
